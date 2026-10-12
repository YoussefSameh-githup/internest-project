"""Universal (all-faculty) skill extractor: document text → explicit + implicit skills.

Primary: LLM (OpenAI-compatible, two keys with failover) returning a strict JSON profile whose skill
names are normalized onto the Skill taxonomy. Fallback: local regex matcher.
"""
import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import IntegrityError, transaction
from django.utils.text import slugify
from django.utils.translation import gettext as _

from .llm import LLM_ERRORS, chat_json, llm_enabled
from .models import Discipline, Skill, StudentSkill
from .taxonomy import TaxonomyIndex

MAX_TEXT_CHARS = 200_000
LLM_INPUT_CHARS = 14_000
LLM_MAX_SKILLS = 20

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    pass


_CID = re.compile(r"\(cid:\d+\)")  # unmapped glyphs from PDF fonts
_CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f​-‏‪-‮﻿�]")
_BULLETS = re.compile(r"^[\s•·▪●○■□◆◇►▶✓✔➢➤*-]+")


def sanitize_text(text: str) -> str:
    """Clean raw PDF/DOCX text: Unicode forms, broken glyphs, hyphenated line breaks, bullets, junk lines."""
    text = unicodedata.normalize("NFKC", text or "")  # ligatures (ﬁ) and Arabic presentation forms
    text = _CONTROL.sub("", _CID.sub(" ", text)).replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"(?<=[a-z])-\n(?=[a-z])", "", text)  # "manage-\nment" → "management"
    lines = []
    for raw in text.split("\n"):
        line = re.sub(r" {2,}", " ", _BULLETS.sub("", raw)).strip()
        if line and sum(ch.isalnum() for ch in line) >= len(line) * 0.2:  # drop lines of mostly symbols
            lines.append(line)
    return "\n".join(lines)


def read_document_text(uploaded_file) -> str:
    name = (uploaded_file.name or "").lower()
    data = uploaded_file.read()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:15])
        elif name.endswith(".docx"):
            import docx
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            text = "\n".join(parts)
        else:
            raise ExtractionError(_("Unsupported file type. Upload a PDF or DOCX."))
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(_("Could not read this file. Make sure it is a valid, non-encrypted PDF or DOCX.")) from exc
    text = sanitize_text(text)[:MAX_TEXT_CHARS]
    if len(text.strip()) < 40:
        raise ExtractionError(_("No readable text found (scanned image PDFs are not supported)."))
    return text


def _normalize(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[إأآ]", "ا", text)
    return text.replace("ى", "ي").replace("ة", "ه")


_ARABIC_PROCLITICS = r"(?:وال|بال|فال|كال|لل|ال|و|ف|ب|ل|ك)?"
_ARABIC_BLOCK = re.compile("[\\u0600-\\u06FF]")


def _term_pattern(term: str) -> re.Pattern:
    words = _normalize(term).split()
    body = r"\s+".join(re.escape(w) for w in words)
    # Arabic attaches conjunctions/prepositions to the next word (e.g. "وصياغة").
    prefix = _ARABIC_PROCLITICS if words and _ARABIC_BLOCK.match(words[0][0]) else ""
    # Lookarounds instead of \b so terms like "c++", "c#", ".net" still match.
    return re.compile(rf"(?<!\w){prefix}{body}(?!\w)")


def _evidence(text: str, start: int, end: int) -> str:
    left = max(text.rfind("\n", 0, start), text.rfind(". ", 0, start)) + 1
    right_candidates = [i for i in (text.find("\n", end), text.find(". ", end)) if i != -1]
    right = min(right_candidates) if right_candidates else len(text)
    snippet = text[left:right].strip()
    return snippet[:280]


@dataclass
class ExtractedSkill:
    skill: Skill
    source: str
    evidence: str


@dataclass
class CVAnalysis:
    skills: list
    candidate_name: str = ""
    summary: str = ""
    years_by_skill: dict = field(default_factory=dict)  # Skill pk → Decimal years
    top_domains: list = field(default_factory=list)
    ai_used: bool = False
    ai_failed: bool = False  # the LLM is configured but every key failed: results come from the local matcher


def analyze_cv(text: str, skills=None) -> CVAnalysis:
    """LLM analysis when configured; if every key fails, the local matcher (flagged with ai_failed)."""
    skills = list(skills if skills is not None else Skill.objects.filter(is_active=True))
    if llm_enabled():
        try:
            with transaction.atomic():  # savepoint: a failed write never poisons the request's transaction
                return analyze_cv_llm(text, skills)
        except LLM_ERRORS as exc:
            logger.warning("LLM skill extraction failed (%s); using local matcher", type(exc).__name__)
            return CVAnalysis(extract_skills_local(text, skills), ai_failed=True)
    return CVAnalysis(extract_skills_local(text, skills))


def extract_skills(text: str, skills=None) -> list[ExtractedSkill]:
    return analyze_cv(text, skills).skills


_DOMAINS = [d.value for d in Discipline]
_SYSTEM_PROMPT = (
    "You are a CV parser. Read the WHOLE CV (summary, experience, projects, education, courses, activities; "
    "English/Arabic, any discipline) and list ALL professional skills, max 20, most relevant first, including "
    "implicit ones shown by projects/work. "
    "Every skill name MUST be a standard taxonomy name: if it is the same skill as a name in Known, return that "
    "Known name EXACTLY (e.g. Python3, py, python developer -> Python); otherwise use the global ESCO / O*NET "
    "preferred label (1-4 words, English, no versions or job titles). "
    "Reply with JSON only, exactly this schema: "
    '{"candidate_name":"","summary":"one sentence, max 40 words","extracted_skills":["Skill"],'
    '"implied_skills":["the extracted_skills that are only implied, not written"],'
    '"years_experience_by_skill":{"Skill":1.5},'
    f'"top_domains":["1-3 of: {", ".join(_DOMAINS)}"]}}'
)


def _get_or_create_skill(name: str, discipline: str, index: TaxonomyIndex):
    skill = index.match(name)
    if skill is None:
        name = index.canonical_name(name)
        if not (2 <= len(name) <= 60) or len(name.split()) > 5:
            return None
        skill = Skill.objects.filter(name__iexact=name).first()  # e.g. an inactive skill
    if skill is None:
        base = slugify(name, allow_unicode=True)[:90] or "skill"
        slug, n = base, 2
        while Skill.objects.filter(slug=slug).exists():
            slug, n = f"{base}-{n}", n + 1
        try:
            skill = Skill.objects.create(name=name, slug=slug, discipline=discipline, aliases=name)
        except IntegrityError:
            skill = Skill.objects.filter(name__iexact=name).first()
    if skill is not None:
        index.add(skill)
    return skill if skill is not None and skill.is_active else None


def _clean_str(value, limit):
    return re.sub(r"\s+", " ", value).strip()[:limit] if isinstance(value, str) else ""


def _years(value):
    try:
        years = Decimal(str(value)).quantize(Decimal("0.1"))
    except (InvalidOperation, ValueError, TypeError):
        return None
    return years if years.is_finite() and Decimal("0") < years <= Decimal("60") else None


_NOISE_LINE = re.compile(r"^(?:[\w.+-]+@[\w-]+\.[\w.]+|\+?[\d\s()./-]{7,}|(?:https?://|www\.)\S+|page \d+(?: of \d+)?)$", re.I)


def compact_cv_text(text: str) -> str:
    """Keep every section but drop tokens that carry no skill signal (contacts, URLs, page footers, repeats)."""
    seen, lines = set(), []
    for raw in sanitize_text(text).splitlines():
        line = re.sub(r"\s+", " ", raw).strip(" \t•·-–—|")
        key = line.lower()
        if len(line) < 2 or key in seen or _NOISE_LINE.match(line):
            continue
        seen.add(key)
        lines.append(line)
    return "\n".join(lines)


def analyze_cv_llm(text: str, skills) -> CVAnalysis:
    cv = compact_cv_text(text)[:LLM_INPUT_CHARS]
    known = ", ".join(s.name for s in skills)
    data = chat_json(_SYSTEM_PROMPT, f"Known: {known}\nCV:\n{cv}", max_tokens=900)
    names = data.get("extracted_skills")
    if not isinstance(names, list):
        raise ValueError("LLM response missing 'extracted_skills' list")

    def as_list(key):
        value = data.get(key)
        return [v for v in value if isinstance(v, str)] if isinstance(value, list) else []

    domains = [d for d in as_list("top_domains") if d in _DOMAINS][:3]
    implied = {n.strip().lower() for n in as_list("implied_skills")}
    raw_years = data.get("years_experience_by_skill")
    years_by_name = {str(k).strip().lower(): v for k, v in raw_years.items()} if isinstance(raw_years, dict) else {}

    index = TaxonomyIndex(skills)
    local = {e.skill.pk: e for e in extract_skills_local(text, skills)}
    analysis = CVAnalysis([], _clean_str(data.get("candidate_name"), 120), _clean_str(data.get("summary"), 400),
                          top_domains=domains, ai_used=True)
    seen = set()
    for raw in names[:LLM_MAX_SKILLS]:
        name = _clean_str(raw, 100)
        skill = _get_or_create_skill(name, domains[0] if domains else Discipline.GENERAL, index) if name else None
        if skill is None or skill.pk in seen:
            continue
        seen.add(skill.pk)
        source = StudentSkill.SOURCE_IMPLICIT if name.lower() in implied else StudentSkill.SOURCE_EXPLICIT
        evidence = local[skill.pk].evidence if skill.pk in local else ""
        analysis.skills.append(ExtractedSkill(skill, source, evidence))
        years = _years(years_by_name.get(name.lower()))
        if years is not None:
            analysis.years_by_skill[skill.pk] = years
    return analysis


def extract_skills_local(text: str, skills=None) -> list[ExtractedSkill]:
    skills = list(skills if skills is not None else Skill.objects.filter(is_active=True))
    norm = _normalize(text)
    source_text = text if len(text) == len(norm) else norm
    results = []
    for skill in skills:
        found = None
        for term in skill.alias_list:
            m = _term_pattern(term).search(norm)
            if m:
                found = ExtractedSkill(skill, StudentSkill.SOURCE_EXPLICIT, _evidence(source_text, m.start(), m.end()))
                break
        if found is None:
            for phrase in skill.implicit_signal_list:
                m = _term_pattern(phrase).search(norm)
                if m:
                    found = ExtractedSkill(skill, StudentSkill.SOURCE_IMPLICIT, _evidence(source_text, m.start(), m.end()))
                    break
        if found:
            results.append(found)
    return results


@transaction.atomic
def save_claimed_skills(student_profile, extracted: list[ExtractedSkill], years_by_skill=None) -> tuple[int, int]:
    """Upsert as Claimed. Never downgrades a Verified/Lag skill. Returns (created, updated)."""
    created = updated = 0
    years_by_skill = years_by_skill or {}
    for item in extracted:
        years = years_by_skill.get(item.skill.pk)
        obj, was_created = StudentSkill.objects.get_or_create(
            student=student_profile, skill=item.skill,
            defaults={"source": item.source, "evidence": item.evidence, "years_experience": years},
        )
        if was_created:
            created += 1
            continue
        if years is not None and obj.years_experience != years:
            obj.years_experience = years
            obj.save(update_fields=["years_experience", "updated_at"])
        if obj.source == StudentSkill.SOURCE_IMPLICIT and item.source == StudentSkill.SOURCE_EXPLICIT:
            obj.source = item.source
            obj.evidence = item.evidence
            obj.save(update_fields=["source", "evidence", "updated_at"])
            updated += 1
    return created, updated
